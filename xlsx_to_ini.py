import sys
import os
import configparser
import openpyxl as xl  # pip install openpyxl needed!
import re
import datetime

variableExpr = re.compile(r"\~[a-z]+\([^\)]*\)")
parameterExpr = re.compile(r"%ls|%s|%S|%i|%I|%u|%d|%[0-9.]*f|%\.\*f")


def segment_check(orig, targ):  # imported from localization.py
    warningcnt = errorcnt = 0
    if set(variableExpr.findall(orig)) != set(variableExpr.findall(targ)):
        warningcnt = 1
    if parameterExpr.findall(orig) != parameterExpr.findall(targ):
        errorcnt = 1

    return (warningcnt, errorcnt)


def main(args):
    if __name__ != "__main__":
        docs = args
    else:
        smartcat_config = configparser.ConfigParser(
            delimiters="=", strict=True, interpolation=None
        )
        smartcat_config.read("smartcat.ini")
        docs = [
            smartcat_config["sc_ko_m"]["documentId_1"].split(","),
            smartcat_config["sc_ko_m"]["documentId_2"].split(","),
        ]

    log = open("mpull.log", "w")

    globalconfig = configparser.ConfigParser(
        delimiters="=", strict=True, interpolation=None
    )
    globalconfig.read("mconfig.ini")  # load settings ini file
    excludekeywords = list(globalconfig["parse"]["excludekeywords"].split(","))

    origindata, transdata, phdata = {}, {}, {}
    overwrited, nodata, excluded, variableError, parameterError = 0, 0, 0, 0, 0
    warn_notrans = 0  # flag for mpull.log

    for doc in docs:
        if len(docs) > 1:
            print(f"Reading {doc[1]}")
        wd = xl.load_workbook(doc[1] + ".xlsx", read_only=True, data_only=True)
        ws = wd.active
        for rowd in ws.iter_rows(min_row=2):  # starts with 2nd row
            if len(rowd) > 4:
                print(f"Warning: irregular rows detected in {rowd}")

            inkey, insource, intarg = rowd[0].value, rowd[1].value, rowd[2].value
            if inkey in transdata != None:
                log.write(
                    f"Warning: overwrite duplicated keyword in {rowd}({inkey}) : {transdata[inkey]}\n"
                )
                overwrited += 1
            # if insource.startswith('\''): intarg='\''+intarg #FIXME: print(insource,intarg)    # add missed \' at first
            transdata[inkey] = intarg
            # origindata[inkey]=insource

    with open("global_ref.ini", "r", encoding="utf​-8-sig") as f:
        config_string = "[DEFAULT]\n" + f.read()
    origindata = configparser.ConfigParser(
        delimiters="=", strict=True, interpolation=None
    )
    origindata.optionxform = str
    origindata.read_string(config_string)

    with open(
        "PHkeywords.ini", "r", encoding="utf​-8-sig"
    ) as f:  # PH segments but used in game.
        config_string = "[DEFAULT]\n" + f.read()
    phdata = configparser.ConfigParser(delimiters="=", strict=True, interpolation=None)
    phdata.optionxform = str
    phdata.read_string(config_string)

    with open(
        "manualkeywords.ini", "r", encoding="utf​-8-sig"
    ) as f:  # missed at original ref data
        config_string = "[DEFAULT]\n" + f.read()
    mndata = configparser.ConfigParser(delimiters="=", strict=True, interpolation=None)
    mndata.optionxform = str
    mndata.read_string(config_string)

    pullfilename = "global_pull_" + datetime.date.today().strftime("%y%m%d")
    if os.path.exists(pullfilename + ".ini"):
        iternum = 1
        while os.path.exists(f"{pullfilename}_{iternum}.ini"):
            iternum += 1
        pullfilename = pullfilename+"_"+iternum

    with open(pullfilename + ".ini", "w", encoding="utf​-8-sig") as f:
        for keyword in origindata["DEFAULT"]:
            if "﻿" in keyword:
                keyword = keyword.replace("﻿", "")
            if keyword in transdata and transdata[keyword] != None:
                tmp_warning, tmp_error = segment_check(
                    origindata["DEFAULT"][keyword], transdata[keyword]
                )
                if tmp_warning:
                    log.write(
                        f"Warning : Variable Format is not matched at '{keyword}'\n"
                    )
                if tmp_error:
                    log.write(
                        f"Error : Parameter Format is not matched at '{keyword}'\n"
                    )
                variableError += tmp_warning
                parameterError += tmp_error
                f.write(keyword + "=" + transdata[keyword] + "\n")
            else:  # doesn't exist in smartcat data
                # print(f"Warning : No translated data of '{keyword}', write original data instead.")
                if (
                    keyword in phdata["DEFAULT"]
                ):  # check if data exist in PHKeywords.ini
                    f.write(keyword + "=" + phdata["DEFAULT"][keyword] + "\n")
                    continue
                else:
                    f.write(keyword + "=" + origindata["DEFAULT"][keyword] + "\n")
                if any(
                    tmp in origindata["DEFAULT"][keyword] for tmp in excludekeywords
                ):
                    excluded += 1
                    continue
                if warn_notrans:
                    log.write(
                        f"Warning : No translated data of '{keyword}', write original data instead.\n"
                    )
                nodata += 1

        for keyword in mndata["DEFAULT"]:
            if keyword in transdata:
                print(
                    f"manual input keyword '{keyword}' is already exist in translated data"
                )
            else:
                f.write(keyword + "=" + mndata["DEFAULT"][keyword] + "\n")

    decnt = 0
    with open("depreciated_keywords.log", "w") as f:
        for keyword in transdata:
            if keyword not in origindata["DEFAULT"]:
                f.write(
                    f"keyword '{keyword}' is not used at original .ini file anymore.\n"
                )
                decnt += 1

    if overwrited + nodata > 1:
        print(
            f"Merge done with {overwrited} overwritten, {nodata} original data uses, {excluded} placeholders"
        )
    else:
        print(f"Merge successfully done with skipping {excluded} placeholders")
    if decnt > 0:
        print(
            "There are depreciated keywords on translated data. please check depreciated_keywords.log file."
        )
    if variableError > 0:
        print(f"There are {variableError} warnings in segments. check mpull.log file.")
    if parameterError > 0:
        print(f"There are {parameterError} errors in segments. check mpull.log file.")


if __name__ == "__main__":
    sys.exit(main(sys.argv))
